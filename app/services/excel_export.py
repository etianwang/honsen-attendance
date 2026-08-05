import openpyxl
from sqlalchemy.orm import Session

from app.models import AttendanceValue, MonthlyRoster, Team
from app.services.attendance_service import compose_display, days_in_month, get_month_entries
from app.services.roster_service import get_roster
from app.services.stats import pivot_labor_stats, team_employee_worksite_days, team_month_stats


def build_month_export(db: Session, team: Team, year: int, month: int) -> openpyxl.Workbook:
    roster: list[MonthlyRoster] = get_roster(db, team.id, year, month)
    stats_rows = team_month_stats(db, roster, year, month)
    stats_by_emp = {r["employee_id"]: r for r in stats_rows}

    # by_code_days from stats_rows is personal-quota-scoped (nonwork tags
    # only, collapsed across every team the person is on) — worksite-site
    # tallies must come from THIS team's own entries specifically, or a
    # dual-team person's site days would either double count or leak the
    # other team's site into this export.
    employee_ids = [r.employee_id for r in roster]
    site_days = team_employee_worksite_days(db, team.id, employee_ids, year, month)

    value_code_by_id = {v.id: v.code for v in db.query(AttendanceValue).all()}

    nonwork_codes = sorted({code for r in stats_rows for code in r["by_code_days"]})
    worksite_codes = sorted({code for days in site_days.values() for code in days})

    total_days = days_in_month(year, month)
    pairs = [(r.employee_id, team.id) for r in roster]
    month_entries = get_month_entries(db, pairs, year, month)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"{year}年{month}月份"[:31]

    headers = (
        ["序号", "姓名", "加班天数"]
        + [f"{c}天数" for c in nonwork_codes]
        + ["超出天数"]
        + [f"{c}天数" for c in worksite_codes]
        + [str(d) for d in range(1, total_days + 1)]
    )
    ws.append(headers)

    for idx, r in enumerate(roster, start=1):
        stats = stats_by_emp.get(
            r.employee_id, {"by_code_days": {}, "overtime_days": 0.0, "excess_over_quota_days": 0.0}
        )
        row = [idx, r.employee.full_name, stats["overtime_days"]]
        row += [stats["by_code_days"].get(c, 0.0) for c in nonwork_codes]
        row.append(stats["excess_over_quota_days"])
        emp_site_days = site_days.get(r.employee_id, {})
        row += [emp_site_days.get(c, 0.0) for c in worksite_codes]

        entries_for_emp = month_entries.get((r.employee_id, team.id), {})
        for day in range(1, total_days + 1):
            entry = entries_for_emp.get(day)
            if entry is None:
                row.append("")
            else:
                am_code = value_code_by_id.get(entry.am_value_id)
                pm_code = value_code_by_id.get(entry.pm_value_id)
                row.append(compose_display(am_code, pm_code, entry.evening_overtime))
        ws.append(row)

    return wb


def build_labor_stats_export(rows: list[dict]) -> openpyxl.Workbook:
    pivot = pivot_labor_stats(rows)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "工地x专业组人工统计"[:31]
    ws.append(["工地 / 专业组"] + pivot["teams"] + ["合计"])
    for site in pivot["sites"]:
        row_values = [pivot["grid"][site][team] for team in pivot["teams"]]
        ws.append([site] + row_values + [sum(row_values)])
    return wb


def build_person_stats_export(rows: list[dict], category_rows: list[dict], category_codes: list[str]) -> openpyxl.Workbook:
    wb = openpyxl.Workbook()

    ws1 = wb.active
    ws1.title = "个人天数汇总"[:31]
    ws1.append(["姓名"] + [f"{c}天数" for c in category_codes] + ["加班天数", "超出配额天数"])
    for r in category_rows:
        row = [r["name"]] + [r["by_code_days"].get(c, 0.0) for c in category_codes]
        row += [r["overtime_days"], r["excess_over_quota_days"]]
        ws1.append(row)

    ws2 = wb.create_sheet("工地x班组明细"[:31])
    ws2.append(["姓名", "班组", "工地", "天数"])
    for r in rows:
        ws2.append([r["name"], r["team"], r["site"], r["man_days"]])

    return wb
